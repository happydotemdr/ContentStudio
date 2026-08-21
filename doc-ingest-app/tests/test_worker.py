# tests/test_worker.py
from unittest.mock import patch

import pytest

from doc_ingest.config import Config
from doc_ingest import jobs, sync, worker


def _seed_pending_job(conn, tmp_input_root, rel_path="Folder/Notes.txt", content=b"hello world this is real text"):
    (tmp_input_root / "Folder").mkdir(parents=True, exist_ok=True)
    (tmp_input_root / rel_path).write_bytes(content)
    sync.sync_source_files(conn, tmp_input_root)
    jobs.enqueue_pending_jobs(conn)
    return jobs.claim_job(conn, worker_id="w1")


def test_process_job_happy_path_writes_commits_locks_and_indexes(conn, tmp_path):
    input_root = tmp_path / "input"
    input_root.mkdir()
    output_root = tmp_path / "output"
    cfg = Config(input_root=input_root, output_root=output_root)
    job_id = _seed_pending_job(conn, input_root)

    with patch("doc_ingest.lock.apply_readonly_lock") as mock_lock, \
         patch("doc_ingest.lock.verify_locked", return_value=True):
        worker.process_job(conn, job_id, cfg, worker_id="w1")

    job_row = conn.execute("SELECT status FROM conversion_jobs WHERE id = ?", (job_id,)).fetchone()
    assert job_row[0] == "complete"

    conversion = conn.execute(
        "SELECT status, locked_confirmed_at, output_path, conversion_tool FROM conversions WHERE job_id = ?", (job_id,)
    ).fetchone()
    assert conversion[0] == "current"
    assert conversion[1] is not None
    assert conversion[3] == "passthrough"  # .txt bypasses firecrawl entirely

    output_file = output_root / "converted" / conversion[2]
    assert output_file.exists()
    assert "hello world" in output_file.read_text(encoding="utf-8")
    mock_lock.assert_called_once_with(output_file)

    fts_row = conn.execute(
        "SELECT body FROM conversions_fts WHERE conversions_fts MATCH 'hello'"
    ).fetchone()
    assert fts_row is not None


def test_process_job_handles_an_extensionless_pdf_via_sniffed_signature(conn, tmp_path):
    input_root = tmp_path / "input"
    input_root.mkdir()
    output_root = tmp_path / "output"
    cfg = Config(input_root=input_root, output_root=output_root)
    job_id = _seed_pending_job(conn, input_root, rel_path="Folder/report", content=b"%PDF-1.4 fake pdf bytes")

    from unittest.mock import MagicMock

    mock_client = MagicMock()
    mock_client.parse.return_value = MagicMock(markdown="# Parsed from a sniffed PDF\n\nreal words here")
    with patch("firecrawl.Firecrawl", return_value=mock_client), \
         patch("doc_ingest.metadata_readers.read_pdf_page_count", return_value=1), \
         patch("doc_ingest.lock.apply_readonly_lock"), \
         patch("doc_ingest.lock.verify_locked", return_value=True):
        worker.process_job(conn, job_id, cfg, worker_id="w1")

    job_row = conn.execute("SELECT status, failure_reason FROM conversion_jobs WHERE id = ?", (job_id,)).fetchone()
    assert job_row[0] == "complete", job_row[1]
    conversion = conn.execute("SELECT source_type FROM conversions WHERE job_id = ?", (job_id,)).fetchone()
    assert conversion[0] == "pdf"


def test_process_job_marks_failed_on_gauntlet_rejection(conn, tmp_path):
    input_root = tmp_path / "input"
    input_root.mkdir()
    output_root = tmp_path / "output"
    cfg = Config(input_root=input_root, output_root=output_root, size_ratio_floor=0.9)
    job_id = _seed_pending_job(conn, input_root, content=b"a source file with plenty of real bytes in it, much more than the tiny converted output below")

    def _fake_convert(staged_path, source_type, cfg_arg):
        from doc_ingest.convert import ConversionResult
        return ConversionResult(success=True, markdown_body="x", tool="firecrawl-parse", error=None)

    with patch("doc_ingest.worker._convert", side_effect=_fake_convert), \
         patch("doc_ingest.lock.apply_readonly_lock"), \
         patch("doc_ingest.lock.verify_locked", return_value=True):
        worker.process_job(conn, job_id, cfg, worker_id="w1")

    job_row = conn.execute("SELECT status, failure_reason FROM conversion_jobs WHERE id = ?", (job_id,)).fetchone()
    assert job_row[0] == "failed"
    assert job_row[1] == "below_size_ratio_floor"
    assert not (output_root / "converted").exists() or not any((output_root / "converted").rglob("*.md"))


def test_process_job_supersedes_the_prior_current_version(conn, tmp_path):
    input_root = tmp_path / "input"
    input_root.mkdir()
    output_root = tmp_path / "output"
    cfg = Config(input_root=input_root, output_root=output_root)

    with patch("doc_ingest.lock.apply_readonly_lock"), patch("doc_ingest.lock.verify_locked", return_value=True):
        job_id_1 = _seed_pending_job(conn, input_root, content=b"version one text content")
        worker.process_job(conn, job_id_1, cfg, worker_id="w1")

        (input_root / "Folder" / "Notes.txt").write_bytes(b"version two, different text content entirely")
        sync.sync_source_files(conn, input_root)
        jobs.enqueue_pending_jobs(conn)
        job_id_2 = jobs.claim_job(conn, worker_id="w1")
        worker.process_job(conn, job_id_2, cfg, worker_id="w1")

    statuses = conn.execute(
        "SELECT version_number, status FROM conversions ORDER BY version_number"
    ).fetchall()
    assert statuses == [(1, "superseded"), (2, "current")]


def test_process_job_leaves_the_job_at_placing_when_lock_confirmation_fails(conn, tmp_path):
    """verify_locked() returning False (no exception -- icacls "succeeded"
    but the read-back didn't confirm it) must NOT mark the job complete,
    per spec §4 step 9: a conversion with locked_confirmed_at unset is not
    done yet, whether or not an exception was involved."""
    input_root = tmp_path / "input"
    input_root.mkdir()
    output_root = tmp_path / "output"
    cfg = Config(input_root=input_root, output_root=output_root)
    job_id = _seed_pending_job(conn, input_root)

    with patch("doc_ingest.lock.apply_readonly_lock"), patch("doc_ingest.lock.verify_locked", return_value=False):
        worker.process_job(conn, job_id, cfg, worker_id="w1")

    job_row = conn.execute("SELECT status FROM conversion_jobs WHERE id = ?", (job_id,)).fetchone()
    assert job_row[0] == "placing"
    conversion = conn.execute("SELECT locked_confirmed_at FROM conversions WHERE job_id = ?", (job_id,)).fetchone()
    assert conversion[0] is None


def test_process_job_resumes_lock_only_after_a_simulated_crash(conn, tmp_path):
    """A job whose .md was written and committed as 'current' but never
    confirmed locked (process died between step 9(b) and 9(d)) must be
    re-locked on the next pass, not re-converted -- and the job itself must
    move to 'complete' only once resume actually succeeds."""
    input_root = tmp_path / "input"
    input_root.mkdir()
    output_root = tmp_path / "output"
    cfg = Config(input_root=input_root, output_root=output_root)
    job_id = _seed_pending_job(conn, input_root)

    with patch("doc_ingest.lock.apply_readonly_lock", side_effect=RuntimeError("simulated crash")):
        with pytest.raises(RuntimeError):
            worker.process_job(conn, job_id, cfg, worker_id="w1")

    conversion = conn.execute(
        "SELECT status, locked_confirmed_at, output_path FROM conversions WHERE job_id = ?", (job_id,)
    ).fetchone()
    assert conversion[0] == "current"
    assert conversion[1] is None  # written, not yet confirmed locked
    job_row = conn.execute("SELECT status FROM conversion_jobs WHERE id = ?", (job_id,)).fetchone()
    assert job_row[0] == "placing"  # NOT complete -- the lock never confirmed
    output_file = output_root / "converted" / conversion[2]
    assert output_file.exists()  # the write already happened

    with patch("doc_ingest.lock.apply_readonly_lock"), patch("doc_ingest.lock.verify_locked", return_value=True):
        worker.resume_unlocked_conversions(conn, cfg)

    conversion_after = conn.execute(
        "SELECT locked_confirmed_at FROM conversions WHERE job_id = ?", (job_id,)
    ).fetchone()
    assert conversion_after[0] is not None
    job_row_after = conn.execute("SELECT status FROM conversion_jobs WHERE id = ?", (job_id,)).fetchone()
    assert job_row_after[0] == "complete"


def test_process_job_updates_heartbeat_while_converting(conn, tmp_path):
    """Without a running heartbeat thread, heartbeat_at is stamped once at
    claim time and never again -- a slow real conversion would eventually
    look stale to reclaim_stale_jobs (Task 7) and get reclaimed out from
    under its own still-running worker. A short interval keeps this test
    fast (~0.15s) and non-flaky rather than waiting through the real
    30-second default."""
    import time

    input_root = tmp_path / "input"
    input_root.mkdir()
    output_root = tmp_path / "output"
    cfg = Config(input_root=input_root, output_root=output_root, reclaim_heartbeat_interval_s=0.02)
    job_id = _seed_pending_job(conn, input_root)

    claimed_heartbeat = conn.execute(
        "SELECT heartbeat_at FROM conversion_jobs WHERE id = ?", (job_id,)
    ).fetchone()[0]

    observed = {}
    real_convert = worker._convert

    def _slow_convert(staged_path, source_type, cfg_arg):
        time.sleep(0.15)  # several heartbeat ticks at the 0.02s interval above
        observed["mid_run"] = conn.execute(
            "SELECT heartbeat_at FROM conversion_jobs WHERE id = ?", (job_id,)
        ).fetchone()[0]
        return real_convert(staged_path, source_type, cfg_arg)

    with patch("doc_ingest.worker._convert", side_effect=_slow_convert), \
         patch("doc_ingest.lock.apply_readonly_lock"), \
         patch("doc_ingest.lock.verify_locked", return_value=True):
        worker.process_job(conn, job_id, cfg, worker_id="w1")

    assert observed["mid_run"] is not None
    assert observed["mid_run"] > claimed_heartbeat  # a heartbeat tick landed on the heartbeat thread's own connection during the slow step
    job_row = conn.execute("SELECT status FROM conversion_jobs WHERE id = ?", (job_id,)).fetchone()
    assert job_row[0] == "complete"


def test_process_job_fails_cleanly_on_a_pre_write_error_instead_of_looping_forever(conn, tmp_path):
    """ANY unexpected exception in the pre-write section (stage/convert/
    gauntlet) must become a clean 'failed' job, not propagate. Left
    uncaught, it would leave the job stuck at 'converting' forever:
    reclaim_stale_jobs would reset it to 'pending', it would be re-claimed,
    and it would crash identically on every future wake -- an unbounded
    poison-pill retry loop, since the job never reaches 'failed' and
    enqueue_pending_jobs's already-failed-at-this-version guard can never
    engage.

    Before Task 22 this was triggered by a .gdoc row hitting
    _source_type_for's KeyError (no Drive branch existed yet). Task 22's
    Drive branch removed that specific trigger, so the failure is injected
    explicitly here via drive_service_factory -- which also models the real
    production case this handler now has to cover: a Drive-native job on a
    machine whose token.json is missing or expired. Injecting it rather
    than relying on token.json's real absence keeps the test deterministic
    on an operator machine that has completed SETUP.md step 6."""
    input_root = tmp_path / "input"
    input_root.mkdir()
    output_root = tmp_path / "output"
    cfg = Config(input_root=input_root, output_root=output_root)
    job_id = _seed_pending_job(conn, input_root, rel_path="Folder/Doc.gdoc", content=b"fake gdoc stub content")

    def _no_drive_credentials(cfg_arg):
        raise RuntimeError("doc-ingest-app has no cached Drive token")

    worker.process_job(conn, job_id, cfg, worker_id="w1", drive_service_factory=_no_drive_credentials)

    job_row = conn.execute("SELECT status, failure_reason FROM conversion_jobs WHERE id = ?", (job_id,)).fetchone()
    assert job_row[0] == "failed"
    assert job_row[1].startswith("unexpected_error:")
    assert not (output_root / "converted").exists() or not any((output_root / "converted").rglob("*.md"))


def test_process_job_resumes_lock_only_after_a_simulated_crash_unchanged_still_propagates(conn, tmp_path):
    """Guard against the pre-write exception handling added above ever
    creeping into the post-write section: an exception from
    lock.apply_readonly_lock must still propagate out of process_job
    uncaught, exactly as before -- that's what leaves the job correctly
    parked at 'placing' (conversion already written+committed) for
    resume_unlocked_conversions, rather than the pre-write handler
    misclassifying it as a job that never got anywhere."""
    input_root = tmp_path / "input"
    input_root.mkdir()
    output_root = tmp_path / "output"
    cfg = Config(input_root=input_root, output_root=output_root)
    job_id = _seed_pending_job(conn, input_root)

    with patch("doc_ingest.lock.apply_readonly_lock", side_effect=RuntimeError("simulated crash")):
        with pytest.raises(RuntimeError):
            worker.process_job(conn, job_id, cfg, worker_id="w1")

    conversion = conn.execute(
        "SELECT status, locked_confirmed_at FROM conversions WHERE job_id = ?", (job_id,)
    ).fetchone()
    assert conversion[0] == "current"
    assert conversion[1] is None
    job_row = conn.execute("SELECT status FROM conversion_jobs WHERE id = ?", (job_id,)).fetchone()
    assert job_row[0] == "placing"


def test_resume_unlocked_conversions_isolates_a_failing_row_and_still_resumes_the_rest(conn, tmp_path):
    """resume_unlocked_conversions IS the crash-recovery mechanism -- one
    row whose icacls call fails (lock.apply_readonly_lock uses
    subprocess.run(..., check=True) and raises CalledProcessError on any
    icacls failure) must not abort the whole sweep and strand every OTHER
    unlocked conversion in the same batch."""
    input_root = tmp_path / "input"
    input_root.mkdir()
    output_root = tmp_path / "output"
    cfg = Config(input_root=input_root, output_root=output_root)

    with patch("doc_ingest.lock.apply_readonly_lock", side_effect=RuntimeError("simulated crash")):
        job_id_1 = _seed_pending_job(conn, input_root, rel_path="Folder/One.txt", content=b"first file content")
        with pytest.raises(RuntimeError):
            worker.process_job(conn, job_id_1, cfg, worker_id="w1")

        job_id_2 = _seed_pending_job(conn, input_root, rel_path="Folder/Two.txt", content=b"second file content")
        with pytest.raises(RuntimeError):
            worker.process_job(conn, job_id_2, cfg, worker_id="w1")

    conversion_1_id, output_path_1 = conn.execute(
        "SELECT id, output_path FROM conversions WHERE job_id = ?", (job_id_1,)
    ).fetchone()
    conversion_2_id, output_path_2 = conn.execute(
        "SELECT id, output_path FROM conversions WHERE job_id = ?", (job_id_2,)
    ).fetchone()
    assert conversion_1_id != conversion_2_id
    assert output_path_1 != output_path_2

    # Compare resolved Path objects, not raw strings -- output_path is
    # stored with forward slashes but Path normalizes to backslashes on
    # Windows, so a substring check against str(path) would never match.
    failing_final_path = cfg.converted_root / output_path_1

    def _apply_lock_first_row_fails(path):
        if path == failing_final_path:
            raise RuntimeError("icacls failed for this one file")

    with patch("doc_ingest.lock.apply_readonly_lock", side_effect=_apply_lock_first_row_fails), \
         patch("doc_ingest.lock.verify_locked", return_value=True):
        resumed = worker.resume_unlocked_conversions(conn, cfg)

    assert resumed == [conversion_2_id]

    conversion_1_locked = conn.execute(
        "SELECT locked_confirmed_at FROM conversions WHERE id = ?", (conversion_1_id,)
    ).fetchone()[0]
    conversion_2_locked = conn.execute(
        "SELECT locked_confirmed_at FROM conversions WHERE id = ?", (conversion_2_id,)
    ).fetchone()[0]
    assert conversion_1_locked is None  # the failing row stays unlocked, not silently dropped
    assert conversion_2_locked is not None

    job_1_status = conn.execute("SELECT status FROM conversion_jobs WHERE id = ?", (job_id_1,)).fetchone()[0]
    job_2_status = conn.execute("SELECT status FROM conversion_jobs WHERE id = ?", (job_id_2,)).fetchone()[0]
    assert job_1_status == "placing"  # unchanged -- never got to 'complete'
    assert job_2_status == "complete"

    event_row = conn.execute(
        "SELECT event_type, conversion_id FROM events WHERE event_type = 'resume_lock_failed'"
    ).fetchone()
    assert event_row is not None
    assert event_row[1] == conversion_1_id


def test_process_job_handles_a_gdoc_via_mocked_drive_export(conn, tmp_path):
    from unittest.mock import MagicMock

    from doc_ingest import frontmatter as frontmatter_mod
    from doc_ingest.config import Config

    input_root = tmp_path / "input"
    input_root.mkdir()
    output_root = tmp_path / "output"
    cfg = Config(input_root=input_root, output_root=output_root)

    stub_path = input_root / "Session Notes.gdoc"
    stub_path.write_text('{"doc_id": "doc-1", "resource_key": "rk1", "email": "admin@freedom2beu.com"}', encoding="utf-8")
    sync.sync_source_files(conn, input_root)
    # Simulates drive_sync.sync_drive_metadata (Task 22) having already
    # populated this -- exercised directly by test_drive_sync.py; this test
    # only needs the column populated to prove process_job uses it, not the
    # full sync flow again.
    conn.execute("UPDATE source_files SET drive_modified_time = '2026-08-10T00:00:00Z' WHERE rel_path = 'Session Notes.gdoc'")
    conn.commit()
    jobs.enqueue_pending_jobs(conn)
    job_id = jobs.claim_job(conn, worker_id="w1")

    def _fake_export_google_doc(service, doc_id, dest_path, cfg_arg):
        from doc_ingest.convert import ConversionResult
        dest_path.write_bytes(b"# Exported directly as markdown\n\nplenty of real words here")
        return ConversionResult(success=True, markdown_body="# Exported directly as markdown\n\nplenty of real words here", tool="google-docs-export", error=None)

    mock_service_factory = lambda cfg_arg: MagicMock()
    with patch("doc_ingest.drive_client.export_google_doc", side_effect=_fake_export_google_doc), \
         patch("doc_ingest.lock.apply_readonly_lock"), \
         patch("doc_ingest.lock.verify_locked", return_value=True):
        worker.process_job(conn, job_id, cfg, worker_id="w1", drive_service_factory=mock_service_factory)

    job_row = conn.execute("SELECT status, failure_reason FROM conversion_jobs WHERE id = ?", (job_id,)).fetchone()
    assert job_row[0] == "complete", job_row[1]
    conversion = conn.execute(
        "SELECT source_type, conversion_tool, drive_modified_time_at_conversion, output_path "
        "FROM conversions WHERE job_id = ?", (job_id,)
    ).fetchone()
    assert conversion[:3] == ("gdoc", "google-docs-export", "2026-08-10T00:00:00Z")

    # source_modified_at must be the Drive edit time, NOT the static local
    # stub's own filesystem mtime (spec §4 step 3, §7).
    output_file = output_root / "converted" / conversion[3]
    fm, _ = frontmatter_mod.parse(output_file.read_text(encoding="utf-8"))
    assert fm["source_modified_at"] == "2026-08-10T00:00:00Z"


def test_process_job_handles_a_gdoc_docx_fallback_export(conn, tmp_path):
    """Proves the docx-fallback filename fix: export_google_doc writes docx
    bytes to a path initially named export.md (it doesn't know the format
    ahead of time); _convert_drive_native must rename it to export.docx
    before treating it as a real docx (independent word-count reader, and
    the file/content-type sent to firecrawl must agree)."""
    from unittest.mock import MagicMock

    import docx as docx_lib

    from doc_ingest.config import Config

    input_root = tmp_path / "input"
    input_root.mkdir()
    output_root = tmp_path / "output"
    cfg = Config(input_root=input_root, output_root=output_root)

    stub_path = input_root / "Long Session Notes.gdoc"
    stub_path.write_text('{"doc_id": "doc-2", "resource_key": "rk2", "email": "admin@freedom2beu.com"}', encoding="utf-8")
    sync.sync_source_files(conn, input_root)
    conn.execute("UPDATE source_files SET drive_modified_time = '2026-08-10T00:00:00Z' WHERE rel_path = 'Long Session Notes.gdoc'")
    conn.commit()
    jobs.enqueue_pending_jobs(conn)
    job_id = jobs.claim_job(conn, worker_id="w1")

    def _fake_export_google_doc(service, doc_id, dest_path, cfg_arg):
        from doc_ingest.convert import ConversionResult
        document = docx_lib.Document()
        document.add_paragraph("word " * 50)
        document.save(dest_path)  # written to a path still named export.md
        return ConversionResult(success=True, markdown_body=None, tool="google-docs-export-docx-fallback", error=None)

    def _fake_convert(staged_path, source_type, cfg_arg):
        from doc_ingest.convert import ConversionResult
        assert staged_path.suffix == ".docx"  # the rename must have already happened
        assert source_type == "docx"
        return ConversionResult(success=True, markdown_body="word " * 50, tool="firecrawl-parse", error=None)

    mock_service_factory = lambda cfg_arg: MagicMock()
    with patch("doc_ingest.drive_client.export_google_doc", side_effect=_fake_export_google_doc), \
         patch("doc_ingest.worker._convert", side_effect=_fake_convert), \
         patch("doc_ingest.lock.apply_readonly_lock"), \
         patch("doc_ingest.lock.verify_locked", return_value=True):
        worker.process_job(conn, job_id, cfg, worker_id="w1", drive_service_factory=mock_service_factory)

    job_row = conn.execute("SELECT status, failure_reason FROM conversion_jobs WHERE id = ?", (job_id,)).fetchone()
    assert job_row[0] == "complete", job_row[1]
    conversion = conn.execute(
        "SELECT source_type, conversion_tool, word_count FROM conversions WHERE job_id = ?", (job_id,)
    ).fetchone()
    assert conversion == ("gdoc", "google-docs-export-docx-fallback", 50)


def test_process_job_handles_a_gsheet_via_mocked_drive_export(conn, tmp_path):
    """The .gsheet half of _convert_drive_native -- the other ~half of the
    real Drive corpus. Distinct from the .gdoc paths above in three ways
    worth their own coverage: source_type resolves to 'gsheet' (not 'gdoc')
    from the extension, the export always goes to export.xlsx with no
    format-fallback rename, and the independent metadata is the
    sheet/row-count pair Gate 1's xlsx/gsheet branch checks rather than a
    word count."""
    from unittest.mock import MagicMock

    import openpyxl

    from doc_ingest.config import Config

    input_root = tmp_path / "input"
    input_root.mkdir()
    output_root = tmp_path / "output"
    cfg = Config(input_root=input_root, output_root=output_root)

    stub_path = input_root / "Budget.gsheet"
    stub_path.write_text('{"doc_id": "sheet-1", "resource_key": "rk3", "email": "admin@freedom2beu.com"}', encoding="utf-8")
    sync.sync_source_files(conn, input_root)
    conn.execute("UPDATE source_files SET drive_modified_time = '2026-08-11T00:00:00Z' WHERE rel_path = 'Budget.gsheet'")
    conn.commit()
    jobs.enqueue_pending_jobs(conn)
    job_id = jobs.claim_job(conn, worker_id="w1")

    row_count = 5

    def _fake_export_google_sheet(service, doc_id, dest_path, cfg_arg):
        from doc_ingest.convert import ConversionResult
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for i in range(row_count):
            sheet.append([f"cell-{i}-a", f"cell-{i}-b"])
        workbook.save(dest_path)
        return ConversionResult(success=True, markdown_body=None, tool="google-sheets-export", error=None)

    def _fake_convert(staged_path, source_type, cfg_arg):
        from doc_ingest.convert import ConversionResult
        assert staged_path.suffix == ".xlsx"
        assert source_type == "xlsx"  # exported bytes are converted as a real xlsx...
        # The workbook above has row_count non-empty rows, and a real
        # export's FIRST row is the header -- so it carries row_count - 1
        # DATA rows. Mirror that: one header, one separator, row_count - 1
        # data rows. This emitted row_count data rows until 2026-08-21,
        # modelling a headerless sheet no real Drive export produces.
        body = "| a | b |\n|---|---|\n" + "".join(f"| c{i} | d{i} |\n" for i in range(row_count - 1))
        return ConversionResult(success=True, markdown_body=body, tool="firecrawl-parse", error=None)

    mock_service_factory = lambda cfg_arg: MagicMock()
    with patch("doc_ingest.drive_client.export_google_sheet", side_effect=_fake_export_google_sheet), \
         patch("doc_ingest.worker._convert", side_effect=_fake_convert), \
         patch("doc_ingest.lock.apply_readonly_lock"), \
         patch("doc_ingest.lock.verify_locked", return_value=True):
        worker.process_job(conn, job_id, cfg, worker_id="w1", drive_service_factory=mock_service_factory)

    job_row = conn.execute("SELECT status, failure_reason FROM conversion_jobs WHERE id = ?", (job_id,)).fetchone()
    assert job_row[0] == "complete", job_row[1]
    conversion = conn.execute(
        "SELECT source_type, conversion_tool, sheet_count, row_count_total, "
        "drive_modified_time_at_conversion FROM conversions WHERE job_id = ?", (job_id,)
    ).fetchone()
    # ...but the RECORDED source_type/tool stay Drive-native, not 'xlsx'/'firecrawl-parse'.
    assert conversion == ("gsheet", "google-sheets-export", 1, row_count, "2026-08-11T00:00:00Z")


def test_process_job_tags_a_session_outlines_file_with_its_client(conn, tmp_path):
    from doc_ingest import clients_db, worker
    clients_db.register_client(
        conn, slug="sean", display_name="Sean", primary_email="sean.carl.tinsley@gmail.com",
        session_outlines_dir="Client Session Outlines/Sean", drive_folder_id="folder1",
    )
    input_root = tmp_path / "input"
    (input_root / "Client Session Outlines" / "Sean").mkdir(parents=True)
    source = input_root / "Client Session Outlines" / "Sean" / "note.txt"
    source.write_text("some session content", encoding="utf-8")

    from doc_ingest.config import Config
    cfg = Config(input_root=input_root, output_root=tmp_path / "output")

    conn.execute(
        "INSERT INTO source_files (rel_path, extension, classification, size_bytes, mtime, "
        "content_hash, first_seen_at, last_seen_at) VALUES "
        "('Client Session Outlines/Sean/note.txt', 'txt', 'convertible', 20, 'm', 'h', 'n', 'n')"
    )
    source_file_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
    conn.execute(
        "INSERT INTO conversion_jobs (source_file_id, status, created_at) VALUES (?, 'pending', 'n')",
        (source_file_id,),
    )
    job_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
    conn.commit()

    with patch("doc_ingest.lock.apply_readonly_lock"), patch("doc_ingest.lock.verify_locked", return_value=True):
        worker.process_job(conn, job_id, cfg, "w1", calendar_service_factory=lambda: None)

    row = conn.execute("SELECT client FROM conversions WHERE source_file_id = ?", (source_file_id,)).fetchone()
    assert row[0] == "sean"

    output_files = list((cfg.converted_root / "Client Session Outlines" / "Sean").glob("*.md"))
    assert len(output_files) == 1
    content = output_files[0].read_text(encoding="utf-8")
    assert "client: sean" in content


def test_process_job_does_not_tag_a_non_client_file(conn, tmp_path):
    from doc_ingest import worker
    input_root = tmp_path / "input"
    (input_root / "Offer & Coaching Framework" / "Current finalized documents").mkdir(parents=True)
    source = input_root / "Offer & Coaching Framework" / "Current finalized documents" / "Vision.txt"
    source.write_text("vision content", encoding="utf-8")

    from doc_ingest.config import Config
    cfg = Config(input_root=input_root, output_root=tmp_path / "output")

    conn.execute(
        "INSERT INTO source_files (rel_path, extension, classification, size_bytes, mtime, "
        "content_hash, first_seen_at, last_seen_at) VALUES "
        "('Offer & Coaching Framework/Current finalized documents/Vision.txt', 'txt', 'convertible', "
        "14, 'm', 'h', 'n', 'n')"
    )
    source_file_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
    conn.execute(
        "INSERT INTO conversion_jobs (source_file_id, status, created_at) VALUES (?, 'pending', 'n')",
        (source_file_id,),
    )
    job_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
    conn.commit()

    with patch("doc_ingest.lock.apply_readonly_lock"), patch("doc_ingest.lock.verify_locked", return_value=True):
        worker.process_job(conn, job_id, cfg, "w1", calendar_service_factory=lambda: None)

    row = conn.execute("SELECT client FROM conversions WHERE source_file_id = ?", (source_file_id,)).fetchone()
    assert row[0] is None


def test_process_job_logs_drift_for_an_unlisted_program_doc(conn, tmp_path):
    from doc_ingest import worker
    input_root = tmp_path / "input"
    (input_root / "Offer & Coaching Framework" / "Current finalized documents").mkdir(parents=True)
    source = input_root / "Offer & Coaching Framework" / "Current finalized documents" / "Brand New Doc.txt"
    source.write_text("new doc content", encoding="utf-8")

    from doc_ingest.config import Config
    cfg = Config(input_root=input_root, output_root=tmp_path / "output")

    conn.execute(
        "INSERT INTO source_files (rel_path, extension, classification, size_bytes, mtime, "
        "content_hash, first_seen_at, last_seen_at) VALUES "
        "('Offer & Coaching Framework/Current finalized documents/Brand New Doc.txt', 'txt', "
        "'convertible', 16, 'm', 'h', 'n', 'n')"
    )
    source_file_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
    conn.execute(
        "INSERT INTO conversion_jobs (source_file_id, status, created_at) VALUES (?, 'pending', 'n')",
        (source_file_id,),
    )
    job_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
    conn.commit()

    with patch("doc_ingest.lock.apply_readonly_lock"), patch("doc_ingest.lock.verify_locked", return_value=True):
        worker.process_job(conn, job_id, cfg, "w1", calendar_service_factory=lambda: None)

    event = conn.execute(
        "SELECT event_type FROM events WHERE source_file_id = ? AND event_type = 'program_source_drift'",
        (source_file_id,),
    ).fetchone()
    assert event is not None

    job_status = conn.execute(
        "SELECT status FROM conversion_jobs WHERE id = ?", (job_id,)
    ).fetchone()[0]
    assert job_status == "complete"
    output_files = list(
        (cfg.converted_root / "Offer & Coaching Framework" / "Current finalized documents").glob("*.md")
    )
    assert len(output_files) == 1
